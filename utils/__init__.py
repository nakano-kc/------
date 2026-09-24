# ライブラリ・モジュールの読み込み
# プログラムで必要な機能を使えるようにする
import streamlit as st #-------------Web画面を作る
from datetime import date
import pandas as pd #----------------表データを扱う
import os #--------------------------コンピューター上のファイルやフォルダなどを操作する
import json #------------------------JSON形式のデータを扱う（今回はカテゴリ情報を扱う）
import matplotlib.pyplot as plt #----グラフを作る
import openpyxl #--------------------Excelを扱う
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import streamlit_authenticator as stauth
import yaml
from yaml.loader import SafeLoader

# Excel出力用のデザイン設定
HEADER_COLOR = "2D3748" # ダークグレー（ 家計簿アプリヘッダー）
ACCENT_COLOR = "4A90D9" # ブルー（タイトル・セクション）
ZEBRA_COLOR = "EDF2F7"  # 薄グレー（交互行）
WHITE = "FFFFFF"

def header_style(cell, color=HEADER_COLOR):
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(color=WHITE, bold=True, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def zebra_style(cell, row_idx):
    if row_idx % 2 == 0:
        cell.fill = PatternFill("solid", fgColor=ZEBRA_COLOR)
    cell.alignment = Alignment(horizontal="left", vertical="center")

def money_style(cell, row_idx):
    if row_idx % 2 == 0:
        cell.fill = PatternFill("solid", fgColor=ZEBRA_COLOR)
    cell.number_format = '#,##0'
    cell.alignment = Alignment(horizontal="right", vertical="center")

def thin_border():
    s = Side(style="thin", color="CBD5E0")
    return Border(left=s, right=s, top=s, bottom=s)


def load_categories():
    """カテゴリを読み込む"""
    # 呼び出された瞬間の最新のユーザー名を取得する
    username = st.session_state.get("username", "guest")
    try:
        with open(f'categories_{username}.json', 'r', encoding='utf-8')as f:
            return json.load(f)
    except FileNotFoundError:
        # デフォルト値
        return {
            "収入": ["給与","副業","その他"],
            "支出": ["外食費", "光熱費", "通信費", "交通費", "娯楽費", "その他"]
        }

# カテゴリをJSONファイルに保存する
def save_categories(categories):
    """ユーザーごとのカテゴリーを保存"""
    # 呼び出された瞬間の最新のユーザー名を取得する
    username = st.session_state.get("username", "guest")
    json_file = f"categories_{username}.json"
    with open(json_file, 'w', encoding='utf-8')as f:
        json.dump(categories, f, ensure_ascii=False, indent=2)

# カテゴリを追加してJSONファイルに保存する
# def add_category(categories, category_type, category_name):
#     """カテゴリを追加"""
#     if category_type not in categories:
#         print(f"エラー: {category_type} は存在しません")
#         return False

#     if category_name in categories[category_type]:
#         print(f"{category_name} は既に登録されています")
#         return False

#     categories[category_type].append(category_name)
#     save_categories(categories)
#     print(f"{category_type} に {category_name} を追加しました")
#     return True

# カテゴリの一覧を表示する
# def show_categories(categories):
#     """カテゴリー一覧を表示"""
#     print("\n=== カテゴリ一覧表示 ===")
#     for category_type, items in categories.items():
#         print(f"\n 【{category_type}】")
#         for i, item in enumerate(items, 1):
#             print(f" {i}. {item}")



# 指定した年月の家計簿データをExcelファイルとして作成する(df == "家計簿データ")
def export_excel(df, selected_year, selected_month):

    # Excelブックを作成
    wb = openpyxl.Workbook()

    # ワークシートを取得して名前を設定
    ws1 = wb.active
    ws1.title = f"{selected_year}年{selected_month}月"

    # 日付をExcel表示用の文字列に変換
    monthly_df = df[
        (df["日付"].dt.year == selected_year) &
        (df["日付"].dt.month == selected_month)
    ].copy()
    monthly_df["日付"] = monthly_df["日付"].dt.strftime("%Y/%m/%d")

    # ーータイトル行
    ws1.merge_cells("A1:E1")
    title_cell = ws1["A1"]
    title_cell.value = f"{selected_year}年{selected_month}月 家計簿データ"
    title_cell.fill = PatternFill("solid", fgColor=ACCENT_COLOR)
    title_cell.font = Font(color=WHITE, bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center", vertical= "center")
    ws1.row_dimensions[1].height = 28

    # ーーヘッダー行
    headers = ["日付", "種別", "カテゴリ", "項目", "金額"]
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=2, column=col, value=h)
        header_style(cell)
        cell.border = thin_border()
    ws1.row_dimensions[2].height = 22

    # ーーデータ行
    for row_idx, (_, row) in enumerate(monthly_df.iterrows(), 1):
        values = [row["日付"], row["種別"], row["カテゴリ"], row["項目"], row["金額"]]
        for col, val in enumerate(values, 1):
            cell = ws1.cell(row=row_idx + 2, column=col, value=val)
            cell.border = thin_border()
            if col == 5:
                money_style(cell, row_idx)
            else:
                zebra_style(cell, row_idx)
        ws1.row_dimensions[row_idx + 2].height = 18

    # カテゴリ別集計
    # ーータイトル
    summary_start = len(monthly_df) + 5
    ws1.merge_cells(f"A{summary_start}:C{summary_start}")
    s_title = ws1[f"A{summary_start}"]
    s_title.value = "カテゴリ別集計"
    s_title.fill = PatternFill("solid", fgColor=ACCENT_COLOR)
    s_title.font = Font(color=WHITE, bold=True, size=11)
    s_title.alignment = Alignment(horizontal="center", vertical="center")

    # ーーヘッダー
    s_headers = ["種別", "カテゴリ", "合計金額"]
    for col, h in enumerate(s_headers, 1):
        cell = ws1.cell(row=summary_start + 1, column=col, value=h)
        header_style(cell)
        cell.border = thin_border()

    # ーーデータ集計
    cat_summary = monthly_df.groupby(["種別", "カテゴリ"])["金額"].sum().reset_index()
    for row_idx, (_, row) in enumerate(cat_summary.iterrows(), 1):
        for col, val in enumerate([row["種別"], row["カテゴリ"], row["金額"]],  1):
            cell = ws1.cell(row=summary_start + 1 + row_idx, column= col, value=val)
            cell.border = thin_border()
            if col == 3:
                money_style(cell, row_idx)
            else:
                zebra_style(cell, row_idx)

    # ーー列幅
    for col, width in zip(range(1, 6), [14, 8, 12, 20, 12]):
        ws1.column_dimensions[get_column_letter(col)].width = width

    # ーーExcelファイルをメモリ上に保存して返す
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer